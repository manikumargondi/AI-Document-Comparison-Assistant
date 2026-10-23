"""Generate downloadable DOCX and PDF comparison reports."""

from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)
from reportlab.lib.styles import getSampleStyleSheet


def _format_value(value):
    """Convert structured results into readable text."""
    if isinstance(value, dict):
        return "\n".join(
            f"{key.replace('_', ' ').title()}: {_format_value(item)}"
            for key, item in value.items()
        )
    if isinstance(value, list):
        return "\n".join(f"- {_format_value(item)}" for item in value)
    return str(value)


def _sections(result):
    """Return report sections in a consistent order."""
    return [
        ("Document A Summary", result.get("summary_a", "")),
        ("Document B Summary", result.get("summary_b", "")),
        ("Similarities", result.get("similarities", [])),
        ("Differences", result.get("differences", [])),
        ("Key Details", result.get("key_details", [])),
        ("Conclusion", result.get("conclusion", "")),
        ("Limitations", result.get("limitations", [])),
    ]


def generate_docx_report(result):
    """Generate a DOCX report and return its bytes."""
    document = Document()
    document.add_heading("AI Document Comparison Report", level=0)

    for heading, content in _sections(result):
        document.add_heading(heading, level=1)
        text = _format_value(content).strip()
        document.add_paragraph(text or "No information available.")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def generate_pdf_report(result):
    """Generate a PDF report and return its bytes."""
    output = BytesIO()
    styles = getSampleStyleSheet()
    story = [
        Paragraph("AI Document Comparison Report", styles["Title"]),
        Spacer(1, 16),
    ]

    for heading, content in _sections(result):
        story.append(Paragraph(heading, styles["Heading2"]))
        text = _format_value(content).strip() or "No information available."
        # Escape special XML characters before creating PDF paragraphs.
        from xml.sax.saxutils import escape

        for line in text.splitlines():
            story.append(Paragraph(escape(line), styles["BodyText"]))
            story.append(Spacer(1, 4))
        story.append(Spacer(1, 10))

    SimpleDocTemplate(output, pagesize=letter).build(story)
    return output.getvalue()